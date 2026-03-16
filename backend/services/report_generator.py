"""Generates the final .docx report from approved findings + template."""

from __future__ import annotations

import base64
import io
import json
import logging
import zipfile
from pathlib import Path
from typing import Any

from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

from models.schemas import Finding, PhotoGroup, PhotoMeta, Session, TemplateSchema

logger = logging.getLogger(__name__)


def generate_report_docx(session: Session) -> bytes:
    """Build the filled report as a .docx, using the original template if it's docx."""
    if session.template_mime in (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "application/msword",
    ) and session.template_bytes:
        doc = _fill_docx_template(session)
    else:
        doc = _build_report_from_scratch(session)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def generate_download_zip(session: Session) -> bytes:
    """Package the report, photos, and raw data into a ZIP."""
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        report_bytes = generate_report_docx(session)
        zf.writestr("report.docx", report_bytes)

        for group in session.groups:
            folder = _safe_name(group.label)
            for pid in group.photo_ids:
                photo = session.photos.get(pid)
                if photo:
                    fname = photo.filename or pid
                    if not Path(fname).suffix:
                        fname += _ext_from_mime(photo.mime_type)
                    zf.writestr(
                        f"photos/{folder}/{fname}",
                        base64.b64decode(photo.b64_data),
                    )

        raw = _build_raw_json(session)
        zf.writestr("raw_analysis.json", json.dumps(raw, indent=2))

    return buf.getvalue()


# ---------------------------------------------------------------------------
# Internal helpers
# ---------------------------------------------------------------------------

def _fill_docx_template(session: Session) -> Document:
    """Open the original .docx template and append findings."""
    doc = Document(io.BytesIO(session.template_bytes))
    _append_findings_section(doc, session)
    return doc


def _build_report_from_scratch(session: Session) -> Document:
    """When template is PDF or missing, build a clean report from scratch."""
    doc = Document()

    style = doc.styles["Title"]
    style.font.size = Pt(24)
    doc.add_heading("Site Visit Inspection Report", level=0)

    if session.template_schema and session.template_schema.raw_instructions:
        doc.add_paragraph(
            "Template reference: " + session.template_schema.raw_instructions[:200]
        ).italic = True

    _append_findings_section(doc, session)
    return doc


def _append_findings_section(doc: Document, session: Session) -> None:
    """Add all approved findings with photos to the document."""
    doc.add_heading("Findings", level=1)

    group_map = {g.id: g for g in session.groups}

    for idx, finding in enumerate(session.findings.values(), 1):
        group = group_map.get(finding.group_id)
        label = group.label if group else f"Finding {idx}"

        doc.add_heading(f"{idx}. {label}", level=2)

        if finding.severity:
            p = doc.add_paragraph()
            run = p.add_run(f"Severity: {finding.severity.upper()}")
            run.bold = True

        if finding.observation:
            doc.add_heading("Observation", level=3)
            doc.add_paragraph(finding.observation)

        if finding.potential_cause:
            doc.add_heading("Potential Cause", level=3)
            doc.add_paragraph(finding.potential_cause)

        if finding.recommendation:
            doc.add_heading("Recommendation", level=3)
            doc.add_paragraph(finding.recommendation)

        if group:
            _insert_group_photos(doc, group, session.photos)

        doc.add_page_break()


def _insert_group_photos(
    doc: Document,
    group: PhotoGroup,
    photos: dict[str, PhotoMeta],
) -> None:
    """Insert key photos for a group into the document."""
    doc.add_heading("Photos", level=3)

    for pid in group.key_photo_ids:
        photo = photos.get(pid)
        if not photo:
            continue
        try:
            img_bytes = base64.b64decode(photo.b64_data)
            img_stream = io.BytesIO(img_bytes)
            doc.add_picture(img_stream, width=Inches(5.5))
            last_paragraph = doc.paragraphs[-1]
            last_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            doc.add_paragraph(photo.filename, style="Caption")
        except Exception:
            logger.exception("Failed to insert photo %s", pid)
            doc.add_paragraph(f"[Photo: {photo.filename} – could not be inserted]")


def _build_raw_json(session: Session) -> dict[str, Any]:
    """Build the raw analysis JSON for download."""
    return {
        "groups": [g.model_dump() for g in session.groups],
        "findings": {k: v.model_dump() for k, v in session.findings.items()},
        "photo_analyses": {k: v.model_dump() for k, v in session.photo_analyses.items()},
        "template_schema": session.template_schema.model_dump() if session.template_schema else None,
        "photo_filenames": {
            pid: p.filename for pid, p in session.photos.items()
        },
    }


def _safe_name(name: str) -> str:
    return "".join(c if c.isalnum() or c in " _-" else "_" for c in name).strip()[:60]


def _ext_from_mime(mime: str) -> str:
    mapping = {
        "image/jpeg": ".jpg",
        "image/png": ".png",
        "image/gif": ".gif",
        "image/webp": ".webp",
        "image/heic": ".heic",
    }
    return mapping.get(mime, ".jpg")
