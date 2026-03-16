"""Stage 4: Analyze the uploaded report template to understand its structure."""

from __future__ import annotations

import base64
import io
import logging
from typing import Any

from agents.base import BaseLLMAgent
from models.schemas import TemplateField, TemplateFieldType, TemplateSchema
from services import llm_client

logger = logging.getLogger(__name__)

SYSTEM_PROMPT = (
    "You are an expert at analyzing document templates for architectural inspection "
    "reports. Your job is to understand the structure of a template so that it can be "
    "programmatically filled with site visit findings, photos, and observations."
)

USER_PROMPT_DOCX = """\
Below is the extracted text content and structure of a Word document template used
for site visit inspection reports.

{content}

Analyze this template and respond with ONLY valid JSON (no markdown):
{{
  "sections": ["list of section names/headers in order"],
  "fields": [
    {{
      "name": "field name or placeholder",
      "field_type": "text|image|table|list",
      "description": "what should go here",
      "location_hint": "where in the document"
    }}
  ],
  "findings_section": "name of the section where individual findings/observations go",
  "has_photo_placeholders": true/false,
  "raw_instructions": "Detailed instructions on how to fill this template with inspection data. Include info about the structure, what tables exist, how photos should be inserted, and any formatting requirements."
}}
"""

USER_PROMPT_PDF = """\
These are page images of a PDF template used for site visit inspection reports.
Analyze the visual layout and structure.

Respond with ONLY valid JSON (no markdown):
{{
  "sections": ["list of section names/headers in order"],
  "fields": [
    {{
      "name": "field name or placeholder",
      "field_type": "text|image|table|list",
      "description": "what should go here",
      "location_hint": "where in the document"
    }}
  ],
  "findings_section": "name of the section where individual findings/observations go",
  "has_photo_placeholders": true/false,
  "raw_instructions": "Detailed instructions on how to recreate this template in Word and fill it with inspection data. Include layout, table structures, photo placements, formatting requirements."
}}
"""


class TemplateAnalyzerAgent(BaseLLMAgent):
    name = "template_analyzer"

    async def execute(
        self,
        template_bytes: bytes,
        mime_type: str,
        filename: str,
    ) -> TemplateSchema:
        if mime_type in (
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "application/msword",
        ):
            return await self._analyze_docx(template_bytes)
        elif mime_type == "application/pdf":
            return await self._analyze_pdf(template_bytes)
        else:
            raise ValueError(f"Unsupported template type: {mime_type}")

    async def _analyze_docx(self, raw: bytes) -> TemplateSchema:
        from docx import Document

        doc = Document(io.BytesIO(raw))
        parts: list[str] = []

        for para in doc.paragraphs:
            style = para.style.name if para.style else ""
            text = para.text.strip()
            if text:
                parts.append(f"[{style}] {text}")

        for table in doc.tables:
            parts.append("[TABLE]")
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells]
                parts.append(" | ".join(cells))
            parts.append("[/TABLE]")

        content = "\n".join(parts)
        prompt = USER_PROMPT_DOCX.format(content=content)

        data = await llm_client.chat_json(
            prompt,
            system=SYSTEM_PROMPT,
            model=self.text_model,
        )
        return self._parse_schema(data, fmt="docx")

    async def _analyze_pdf(self, raw: bytes) -> TemplateSchema:
        import fitz  # PyMuPDF

        pdf = fitz.open(stream=raw, filetype="pdf")
        images: list[tuple[str, str]] = []

        for page in pdf:
            pix = page.get_pixmap(dpi=150)
            img_bytes = pix.tobytes("png")
            b64 = base64.b64encode(img_bytes).decode()
            images.append((b64, "image/png"))

        pdf.close()

        data = await llm_client.chat_json_with_images(
            USER_PROMPT_PDF,
            images,
            system=SYSTEM_PROMPT,
            model=self.vision_model,
        )
        return self._parse_schema(data, fmt="pdf")

    @staticmethod
    def _parse_schema(data: dict[str, Any], fmt: str) -> TemplateSchema:
        fields = []
        for f in data.get("fields", []):
            ft = f.get("field_type", "text")
            try:
                field_type = TemplateFieldType(ft)
            except ValueError:
                field_type = TemplateFieldType.TEXT
            fields.append(
                TemplateField(
                    name=f.get("name", ""),
                    field_type=field_type,
                    description=f.get("description", ""),
                    location_hint=f.get("location_hint", ""),
                )
            )
        return TemplateSchema(
            format=fmt,
            sections=data.get("sections", []),
            fields=fields,
            findings_section=data.get("findings_section", ""),
            has_photo_placeholders=data.get("has_photo_placeholders", False),
            raw_instructions=data.get("raw_instructions", ""),
        )
